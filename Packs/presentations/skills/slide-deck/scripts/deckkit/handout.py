"""A printable rehearsal artifact: one page per slide, image above, notes below.

The notes hold the spoken script, the clock and the contingencies, and there was no way
to get them out of the file. This composes a PDF — the slide image, then its notes — which
is what a speaker rehearses from, and what an examiner occasionally asks for as a
handout.

It reuses `render`'s slide images, running the render when they are missing or older than
the deck, so `deck handout` is one command from source to something you can hold.
"""

from __future__ import annotations

import math
from pathlib import Path

from .deck import Deck
from .errors import DeckError
from .render import render

#: A4 portrait, in points.
A4_W, A4_H = 595.0, 842.0
MARGIN = 42.0
GAP = 16.0
MIN_FONT, MAX_FONT = 5.5, 9.5
#: Slide images are downscaled to this width before embedding. Without it the handout
#: carries 32 full-resolution renders and lands around 100 MB.
TARGET_W = 1500

#: A Unicode face for the notes. The built-in font is Latin-1 and renders the German
#: quotes and the en dash as `?`, which is exactly the typography a German deck needs.
_FONTS = (
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
)
#: Fallback only, when no Unicode font is on the machine: the same characters as ASCII.
_ASCII = str.maketrans({"\u2013": "-", "\u2014": "-", "\u201e": '"',
                        "\u201c": '"', "\u201d": '"', "\u2018": "'",
                        "\u2019": "'", "\u2026": "...", "\u2192": "->"})


def build(pptx: str | Path, out_dir: str | Path | None = None,
          dpi: int = 110) -> Path:
    import pymupdf

    pptx = Path(pptx)
    if not pptx.exists():
        raise DeckError(
            f"cannot build a handout for {pptx}: not built yet.\n"
            f"  Run `scripts/deck build <deck.py>` first."
        )
    out_dir = Path(out_dir) if out_dir else pptx.parent / "render"

    images = sorted(out_dir.glob("slide-*.png"))
    if not images or min(p.stat().st_mtime for p in images) < pptx.stat().st_mtime:
        images = render(pptx, out_dir=out_dir, dpi=dpi)["slides"]

    notes = _notes(pptx)
    font = next((p for p in _FONTS if Path(p).exists()), None)
    document = pymupdf.open()
    for index, image in enumerate(images):
        _page(document, image, notes[index] if index < len(notes) else "", font)
    target = out_dir / f"{pptx.stem}-handout.pdf"
    document.save(target)
    document.close()
    return target


def _notes(pptx: Path) -> list[str]:
    from pptx import Presentation

    out = []
    for slide in Presentation(str(pptx)).slides:
        text = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        out.append(text.strip())
    return out


def _page(document, image: Path, note: str, font: str | None) -> None:
    import io

    import pymupdf
    from PIL import Image

    page = document.new_page(width=A4_W, height=A4_H)
    with Image.open(image) as source:
        rgb = source.convert("RGB")
        aspect = rgb.size[0] / rgb.size[1]
        if rgb.size[0] > TARGET_W:
            rgb = rgb.resize((TARGET_W, round(rgb.size[1] * TARGET_W / rgb.size[0])),
                             Image.LANCZOS)
        buffer = io.BytesIO()
        rgb.save(buffer, format="JPEG", quality=85)
    width = A4_W - 2 * MARGIN
    height = width / aspect
    page.insert_image(pymupdf.Rect(MARGIN, MARGIN, MARGIN + width, MARGIN + height),
                      stream=buffer.getvalue())
    if not note:
        return
    text_rect = pymupdf.Rect(MARGIN, MARGIN + height + GAP, A4_W - MARGIN, A4_H - MARGIN)
    size = MAX_FONT
    while size > MIN_FONT and _needed_height(note, text_rect.width, size) > text_rect.height:
        size -= 0.5
    if font:
        page.insert_textbox(text_rect, note, fontsize=size, fontfile=font, fontname="deck")
    else:
        page.insert_textbox(text_rect, note.translate(_ASCII), fontsize=size,
                            fontname="helv", align=0)


def _needed_height(text: str, width: float, size: float) -> float:
    """Estimate the height `text` occupies at `size`. Over-reports, like `check`."""
    per_line = max(width / (size * 0.52), 1)
    lines = sum(max(1, math.ceil(len(line) / per_line)) for line in text.split("\n"))
    return lines * size * 1.25
