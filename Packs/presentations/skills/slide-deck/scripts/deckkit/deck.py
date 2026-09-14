"""The presentation: theme, page counter, slide archetypes, save.

A deck file is a Python module that builds one `Deck` and calls its methods. There
is no separate content format on purpose — a deck's structure is a sequence of
decisions with numbers in it (time budgets, run counts, matrix rows), and Python
reads those better than YAML does.

    from deckkit import Deck, Theme

    deck = Deck(theme=Theme.load("Assets/themes/warm.json"),
                name="my-talk", author="Lars Boes",
                footer_label="Talk / 18.09.2026", output="out")

    deck.title(kicker="...", title="...", meta=["..."])
    deck.agenda(["Problem", "Method", "Results"])

    s = deck.open("An assertion, not a topic", "1 · Problem")
    s.three_columns([...])
    s.notes("20 s. ...")

`scripts/deck build <file>` runs the module and saves whatever `deck` it finds.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from . import layout as L
from .errors import DeckError
from .slide import Slide, plain_text
from .theme import Theme


class Deck:
    def __init__(self, theme: Theme, *, name: str = "deck", author: str = "",
                 footer_label: str = "", output: str | Path = "out",
                 title_text: str | None = None, max_words: int = 120,
                 minutes: float | None = None, export_pdf: bool = False):
        self.theme = theme
        self.name = name
        self.author = author
        self.footer_label = footer_label
        self.output = Path(output)
        self.title_text = title_text
        #: Words per slide above which `check` calls the slide dense. Raise it for a
        #: deck whose audience reads the slide while listening — a summary panel at
        #: the end of a section is dense on purpose, and the default nudge would
        #: fire on every one of them until it stopped meaning anything.
        self.max_words = max_words
        #: The talk's allotted length, for `deck timing`. Optional: without it, timing
        #: reports the noted budget and the implied duration and compares nothing.
        self.minutes = minutes
        #: Write the PDF beside the `.pptx` on every `build`, so the deck folder is the
        #: deliverable and the export is not a step to remember. Needs LibreOffice.
        self.export_pdf = export_pdf

        self.prs = Presentation()
        self.prs.slide_width = Inches(theme.metric("width"))
        self.prs.slide_height = Inches(theme.metric("height"))
        # The title slide occupies page 1 and prints no number, so the agenda is 2.
        # Numbering from the physical position keeps "see slide 8" true.
        self._page = 1
        self.slides: list[Slide] = []
        self._toc: list[str] = []
        # Set by backup_divider(); every slide opened afterwards is a backup slide, so
        # `readiness` can split the talk from the appendix without parsing titles.
        self._in_backup = False

    # ── internals ───────────────────────────────────────────────────────
    def _blank(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        fill = slide.background.fill
        fill.solid()
        # White, not the palette's paper tone: figures are rendered on white and a
        # tinted background shows up as a rectangle around every one of them.
        fill.fore_color.rgb = self.theme.rgb("white")
        # Drawn first, so it is the bottom layer and every panel sits on top of it.
        if self.theme.metric("flourish") > 0:
            L.flourish(slide, self.theme)
        return slide

    def _next(self) -> int:
        self._page += 1
        return self._page

    def _slide(self, shape_slide, top: float) -> Slide:
        s = Slide(shape_slide, self.theme, top, self._next(), self.footer_label,
                  self.author)
        self.slides.append(s)
        return s

    def _header(self, shape_slide, title: str, subtitle: str | None):
        theme = self.theme
        lines = title.split("\n")
        line_h = theme.metric("line_h")
        x, w = theme.metric("margin"), theme.content_width
        L.write(shape_slide, theme, x, theme.metric("title_y"), w,
                line_h * len(lines) + 0.1, lines, size=theme.size("title"),
                bold=True, line_spacing=0.92, gap=0)
        y = theme.metric("title_y") + line_h * len(lines)
        if subtitle:
            L.write(shape_slide, theme, x, y + 0.02, w, theme.metric("sub_h"),
                    subtitle, size=theme.size("subtitle"),
                    colour=theme.hex("accent"), line_spacing=0.95)
            y += theme.metric("sub_h")
        # A theme with `title_rule` at 0 carries its assertions in claim bars, so
        # there is no heading for a rule to belong to. Skipping it also returns a
        # tighter top, which is what keeps a titleless slide from starting low.
        if theme.metric("title_rule") > 0:
            rule_y = y + 0.10
            L.rule(shape_slide, theme, x, rule_y, w,
                   thickness=theme.metric("title_rule"))
            return rule_y + theme.metric("gap") + 0.14
        return y + theme.metric("gap") + 0.14

    # ── slides ──────────────────────────────────────────────────────────
    def open(self, title: str = "", subtitle: str | None = None) -> Slide:
        """A content slide: optional header, the accent rule, the footer.

        Passing no `title` gives the claim-first slide: nothing is drawn at the top,
        content starts at `content_y`, and the slide's assertion lives in its claim
        bar (`statement`/`headline`) at the bottom. `check` cannot tell the two
        shapes apart, so `readiness` reads the first claim as the assertion.
        """
        shape_slide = self._blank()
        top = (self._header(shape_slide, title, subtitle) if title
               else self.theme.metric("content_y"))
        slide = self._slide(shape_slide, top)
        slide.title = title
        slide.kind = "backup" if self._in_backup else "content"
        slide._footer()
        return slide

    def title(self, *, kicker: str, title: str, meta: list[str] | None = None,
              date: str = "", badge: str | None = None, tagline: str = "",
              notes: str = "") -> Slide:
        theme = self.theme
        shape_slide = self._blank()
        x, w = theme.metric("margin"), theme.content_width
        L.write(shape_slide, theme, x, 1.62, w, 0.4, kicker,
                size=theme.size("subtitle") + 3.5, bold=True)
        L.write(shape_slide, theme, x, 2.22, w - 1.0, 1.5, title,
                size=theme.size("title") + 1, bold=True, italic=True, line_spacing=1.06)
        if theme.metric("title_rule") > 0:
            L.rule(shape_slide, theme, x, 3.86, w,
                   thickness=theme.metric("title_rule"))
        if date:
            L.write(shape_slide, theme, x, 4.04, 6.0, 0.3, date,
                    size=theme.size("body") + 0.5, colour=theme.hex("accent_mid"))
        if meta:
            L.write(shape_slide, theme, x, 4.92, 8.4, 1.0, meta,
                    size=theme.size("body"), colour=theme.hex("accent_mid"),
                    line_spacing=1.15, gap=2)
        name = badge if badge is not None else self.author
        if name:
            L.badge(shape_slide, theme, x, 6.10, name, height=0.40,
                    size=theme.size("body") + 1)
        if tagline:
            L.write(shape_slide, theme, theme.metric("width") - x - 3.2, 6.16,
                    3.2, 0.3, tagline, size=theme.size("small"),
                    colour=theme.hex("accent_mid"), align=PP_ALIGN.RIGHT)
        slide = Slide(shape_slide, theme, 4.40, 1, self.footer_label, self.author, notes)
        self.slides.append(slide)   # page 1: registered directly, not via _next()
        slide.title = title
        slide.kind = "title"
        if notes:
            slide.notes(notes)
        return slide

    def agenda(self, items: list[str], *, active: int | None = None,
               title: str = "Agenda", subtitle: str | None = None,
               claim: str | None = None, notes: str = "") -> Slide:
        """The contents rail. `active` lights one item; call again per section.

        With `active` and a `claim`, the same call becomes a section divider that keeps
        the running agenda on screen and marks the section the talk has reached. A talk
        that repeats its agenda costs the audience nothing and saves them the question
        of where they are.

        The list is not framed by rules, and the active item is marked by the accent bar
        alone rather than by weight and colour. On a divider the section name is already
        on the slide once, in the heading; repeating it in bold turns the rail into a
        second title and the slide into a stutter. The bar says where the talk is.
        """
        theme = self.theme
        shape_slide = self._blank()
        top = self._header(shape_slide, title, subtitle)
        slide = self._slide(shape_slide, top)
        rail_x, rail_w = 1.95, 9.1
        first_y = 2.10
        # The step adapts to the number of items. A fixed 0.52 fits six entries and
        # runs an eight-entry agenda into the footer, which is exactly the kind of
        # arithmetic that should live here rather than in every deck that grows a
        # section. Seven items or fewer keep the original rhythm; more tighten, and
        # the claim line stays just under the list either way.
        claim_top = 6.00
        step = min(0.52, (claim_top - first_y - 0.16) / max(len(items), 1))
        y = first_y
        for index, item in enumerate(items):
            on = active is not None and index == active
            if on:
                L.rect(shape_slide, theme, rail_x - 0.16, y - 0.03, 0.07,
                       min(0.34, step - 0.16), fill=theme.hex("accent"))
            L.write(shape_slide, theme, rail_x, y, rail_w, min(0.42, step),
                    item, size=theme.size("subtitle") + 2,
                    colour=theme.hex("accent_mid"), bold=False)
            y += step
        if claim:
            L.write(shape_slide, theme, rail_x, y + 0.16, rail_w - 0.6, 0.6, claim,
                    size=theme.size("subtitle") - 1,
                    colour=theme.hex("accent_mid"), line_spacing=1.15)
        slide._footer()
        if notes:
            slide.notes(notes)
        self._toc = list(items)
        slide.title = title
        slide.kind = "agenda"
        return slide

    def section(self, number, title: str, claim: str = "", *, notes: str = "") -> Slide:
        """A divider. `title` may carry a newline; `claim` is one line."""
        theme = self.theme
        shape_slide = self._blank()
        x, w = theme.metric("margin"), theme.content_width
        lines = title.count("\n") + 1
        L.write(shape_slide, theme, x, 2.02, w, 0.32, f"TEIL {number}",
                size=theme.size("body") - 1, colour=theme.hex("accent_mid"), bold=True)
        L.rule(shape_slide, theme, x, 2.46, 1.5, thickness=0.05)
        L.write(shape_slide, theme, x, 2.64, w, 0.64 * lines, title.split("\n"),
                size=theme.size("section"), bold=True, line_spacing=0.95, gap=0)
        if claim:
            L.write(shape_slide, theme, x, 2.64 + 0.64 * lines + 0.16, w - 2.0, 0.9,
                    claim, size=theme.size("subtitle") - 2,
                    colour=theme.hex("accent_mid"), line_spacing=1.15)
        slide = self._slide(shape_slide, 2.02)
        slide.title = title
        slide.kind = "section"
        if claim:
            slide.claims.append(plain_text(claim))
        slide._footer()
        if notes:
            slide.notes(notes)
        return slide

    def backup_divider(self, title: str = "Backup", claim: str = "", notes: str = "") -> Slide:
        theme = self.theme
        shape_slide = self._blank()
        x, w = theme.metric("margin"), theme.content_width
        L.rule(shape_slide, theme, x, 2.50, 1.4, thickness=0.05)
        L.write(shape_slide, theme, x, 2.60, w, 1.2, title,
                size=theme.size("section") + 6, bold=True)
        if claim:
            L.write(shape_slide, theme, x, 3.72, 8.5, 1.0, claim,
                    size=theme.size("statement"), colour=theme.hex("accent_mid"))
        slide = Slide(shape_slide, theme, 2.60, "B0", self.footer_label, self.author)
        self.slides.append(slide)
        self._in_backup = True
        slide.title = title
        slide.kind = "backup_divider"
        if claim:
            slide.claims.append(plain_text(claim))
        slide._footer()
        if notes:
            slide.notes(notes)
        return slide

    def closing(self, *, kicker: str, verdict: str | list[str], thanks: str,
                meta: str = "", answers: list[tuple[str, str]] | None = None,
                notes: str = "") -> Slide:
        """The last slide: the answer, in the shape of the question.

        `answers` is [(label, outcome), ...] drawn as a row of panels above the
        verdict. A talk whose research question has parts ends better with those
        parts answered than with one paragraph that answers them implicitly: the
        closing is where the room checks whether every part was addressed, and a
        list answers that at a glance.

        With `answers` the block sits higher and the verdict is set smaller, because
        the panels are now the content and the verdict is its summary. Without it the
        original, roomier geometry is kept, so an existing deck does not move.

        `verdict` may be a list of paragraphs; a one-sentence answer and its bound are
        two paragraphs, not one long one.
        """
        theme = self.theme
        shape_slide = self._blank()
        x, w = theme.metric("margin"), theme.content_width
        verdict_text = " ".join(verdict) if isinstance(verdict, list) else verdict

        if answers is None:
            L.write(shape_slide, theme, x, 1.55, w, 0.6, kicker,
                    size=theme.size("statement") - 0.5,
                    colour=theme.hex("accent_mid"), bold=True)
            L.rule(shape_slide, theme, x, 2.05, 2.6, thickness=0.045)
            L.write(shape_slide, theme, x, 2.45, w - 0.5, 2.4, verdict,
                    size=theme.size("section") - 12, line_spacing=1.18)
            thanks_y, rule_y = 4.86, 5.62
        else:
            L.write(shape_slide, theme, x, 1.10, w, 0.6, kicker,
                    size=theme.size("statement") - 0.5,
                    colour=theme.hex("accent_mid"), bold=True)
            L.rule(shape_slide, theme, x, 1.60, 2.6, thickness=0.045)
            gap = 0.34
            width = theme.span(len(answers), gap)
            for index, (label, outcome) in enumerate(answers):
                L.panel(shape_slide, theme, x + index * (width + gap), 1.95,
                        width, 1.35, [f"**{label}**", outcome],
                        voice="accent", size=theme.size("body") - 0.5,
                        gap=5, pad=0.20, anchor=MSO_ANCHOR.MIDDLE)
            L.write(shape_slide, theme, x, 3.60, w - 0.5, 1.55, verdict,
                    size=theme.size("section") - 18, line_spacing=1.15)
            thanks_y, rule_y = 5.45, 5.95

        L.write(shape_slide, theme, x, thanks_y, w, 0.5, thanks,
                size=theme.size("subtitle"), colour=theme.hex("accent"), bold=True)
        L.rule(shape_slide, theme, x, rule_y, w, colour=theme.hex("hairline"),
               thickness=0.012)
        if meta:
            L.write(shape_slide, theme, x, rule_y + 0.18, 7.0, 0.3, meta,
                    size=theme.size("small"), colour=theme.hex("accent_mid"))
        slide = self._slide(shape_slide, 2.45 if answers is None else 1.95)
        slide.title = kicker
        slide.kind = "closing"
        slide.claims.append(plain_text(verdict_text))
        slide._footer()
        if notes:
            slide.notes(notes)
        return slide

    # ── output ──────────────────────────────────────────────────────────
    def save(self, path: str | Path | None = None) -> Path:
        if not self.slides:
            raise DeckError("this deck has no slides; call title()/open() first.")
        target = Path(path) if path else self.output / f"{self.name}.pptx"
        target.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(target)
        return target

    def __repr__(self) -> str:
        return f"<Deck {self.name}: {len(self.slides)} slides, theme={self.theme.name}>"

    def argument(self) -> list[dict]:
        """The argument skeleton, in slide order, for `deck readiness`.

        Each entry is `{number, kind, title, claims, notes}`. It is the thing `check`
        and `render` cannot see: whether the deck answers the question it was built
        for. `readiness` prints it; the five-question review is the reviewer's.
        """
        return [{"number": s.number, "kind": s.kind, "title": s.title,
                 "claims": list(s.claims), "notes": s._notes}
                for s in self.slides]
